"""CSV/엑셀 명단 읽기·저장, 대진표 엑셀 저장."""
import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from models import Player, make_player
from scheduler import KIND_EVENT, Schedule

COLUMNS = ["이름", "나이", "성별", "급수", "고정파트너"]
# 머리글 이름이 조금 달라도 알아듣게
ALIASES = {
    "이름": ["이름", "성명", "name"],
    "나이": ["나이", "age"],
    "성별": ["성별", "gender", "sex"],
    "급수": ["급수", "grade", "level"],
    "고정파트너": ["고정파트너", "고정 파트너", "파트너", "partner"],
}


# ---------------------------------------------------------------------------
# 명단 읽기 / 저장
# ---------------------------------------------------------------------------

def _rows_to_players(rows: list[list]) -> list[Player]:
    if not rows:
        raise ValueError("파일이 비어 있어요")
    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    col = {}
    for key, names in ALIASES.items():
        for i, h in enumerate(header):
            if h in [n.lower() for n in names]:
                col[key] = i
                break
    missing = [k for k in COLUMNS[:4] if k not in col]
    if missing:
        raise ValueError(f"머리글에 {', '.join(missing)} 열이 없어요 (필요: {', '.join(COLUMNS)})")

    players, errors = [], []
    for line, row in enumerate(rows[1:], start=2):
        def get(key):
            i = col.get(key)
            if i is None or i >= len(row) or row[i] is None:
                return ""
            return str(row[i]).strip()
        if not any(str(v).strip() for v in row if v is not None):
            continue
        try:
            players.append(make_player(get("이름"), get("나이"), get("성별"),
                                       get("급수"), get("고정파트너")))
        except ValueError as e:
            errors.append(f"{line}행: {e}")
    if errors:
        raise ValueError("\n".join(errors))
    return players


def load_players(path) -> list[Player]:
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        return _rows_to_players(rows)
    # 한국어 엑셀에서 저장한 CSV는 cp949인 경우가 많아요
    for enc in ("utf-8-sig", "cp949"):
        try:
            with open(path, encoding=enc, newline="") as f:
                return _rows_to_players(list(csv.reader(f)))
        except UnicodeDecodeError:
            continue
    raise ValueError("파일 인코딩을 읽을 수 없어요 (UTF-8 또는 CP949)")


def save_players(players: list[Player], path) -> None:
    path = Path(path)
    rows = [[p.name, p.age, p.gender, p.grade, p.partner] for p in players]
    if path.suffix.lower() == ".csv":
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(COLUMNS)
            w.writerows(rows)
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "참가자 명단"
    ws.append(COLUMNS)
    for r in rows:
        ws.append(r)
    _style_table(ws, len(COLUMNS), len(rows) + 1)
    wb.save(path)


# ---------------------------------------------------------------------------
# 대진표 저장
# ---------------------------------------------------------------------------

THIN = Side(style="thin", color="999999")
THICK = Side(style="medium", color="333333")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
EVENT_FILL = PatternFill("solid", fgColor="FCE4D6")
REVIEW_FILL = PatternFill("solid", fgColor="FDE4E1")
ROUND_FILLS = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="EEF3FA")]
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_table(ws, n_cols, n_rows, widths=None):
    for c in range(1, n_cols + 1):
        cell = ws.cell(1, c)
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, CENTER
    for r in range(1, n_rows + 1):
        for c in range(1, n_cols + 1):
            cell = ws.cell(r, c)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if r > 1:
                cell.alignment = CENTER
    for c in range(1, n_cols + 1):
        width = widths[c - 1] if widths else 12
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"


def team_text(team):
    return " · ".join(f"{p.name}({p.grade})" for p in team)


def review_rows(sched: Schedule):
    """검토 필요 경기 목록: (라운드·코트, 대진) 쌍."""
    return [(f"  └ {gm.round}라운드 {gm.court}코트",
             f"{team_text(gm.team_a)}  vs  {team_text(gm.team_b)}") for gm in sched.review_games]


def save_schedule(sched: Schedule, path) -> None:
    wb = Workbook()

    # 1. 대진표
    ws = wb.active
    ws.title = "대진표"
    headers = ["라운드", "코트", "구분", "종류", "팀 A", "", "팀 B", "휴식"]
    rest_col = len(headers)
    ws.append(headers)
    row = 2
    for r, (games, resting) in enumerate(zip(sched.rounds, sched.resting)):
        start = row
        for gm in games:
            ws.append([f"{gm.round}R", gm.court, gm.kind, gm.gender_type,
                       team_text(gm.team_a), "vs", team_text(gm.team_b), ""])
            if gm.needs_review:
                fill = REVIEW_FILL
            else:
                fill = EVENT_FILL if gm.kind == KIND_EVENT else ROUND_FILLS[r % 2]
            for c in range(1, len(headers) + 1):
                ws.cell(row, c).fill = fill
            row += 1
        ws.cell(start, rest_col, ", ".join(p.name for p in resting) or "-")
        if row - 1 > start:
            ws.merge_cells(start_row=start, start_column=1, end_row=row - 1, end_column=1)
            ws.merge_cells(start_row=start, start_column=rest_col, end_row=row - 1, end_column=rest_col)
    _style_table(ws, len(headers), row - 1, [8, 6, 10, 10, 30, 4, 30, 24])
    # 라운드 경계를 굵은 선으로
    row = 2
    for games in sched.rounds:
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row, c)
            cell.border = Border(left=THIN, right=THIN, top=THICK, bottom=THIN)
        row += len(games)
    for c in (5, 7):
        for r in range(2, row):
            ws.cell(r, c).font = Font(bold=True)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0

    # 2. 인원별 통계
    ws = wb.create_sheet("인원별 통계")
    headers = ["이름", "성별", "나이", "급수", "총 게임", "경기 패턴", "휴식",
               "최대 연속", "연속 3게임+", "같은 급수", "이벤트", "남복", "여복", "혼복",
               "혼성 대결", "파트너 중복", "상대 중복"]
    ws.append(headers)
    for s in sorted(sched.stats, key=lambda s: (-s.player.level, s.player.name)):
        p = s.player
        ws.append([p.name, p.gender, p.age, p.grade, s.games, s.pattern, s.rests,
                   s.max_streak, s.three_in_row, s.same_grade, s.event, s.men_doubles,
                   s.women_doubles, s.mixed_doubles, s.cross, s.partner_repeats,
                   s.opponent_repeats])
    n = len(sched.stats) + 1
    _style_table(ws, len(headers), n, [10, 6, 6, 6, 8, max(12, len(sched.rounds) * 2.4),
                                       6, 9, 11, 9, 8, 6, 6, 6, 9, 10, 9])
    ws.append([])
    ws.append(["※ 경기 패턴: ● 경기 / ○ 휴식 (왼쪽부터 1라운드)"])

    # 3. 요약
    ws = wb.create_sheet("요약")
    ws.append(["항목", "값"])
    st = sched.settings
    settings_rows = [
        ("코트 수 (입력)", st.courts),
        ("1인당 게임 수 (입력)", st.games_per_player),
        ("성별 모드", st.gender_mode),
        ("성별 구성이 다른 대결 허용", "예" if st.allow_cross else "아니오"),
        ("이벤트 게임 비율", f"{round(st.event_ratio * 100)}%"),
    ]
    review = []
    for k, v in settings_rows + list(sched.summary.items()):
        ws.append([k, v])
        if k.startswith("검토 필요"):  # 검토 필요 경기 목록을 바로 아래에
            for label, text in review_rows(sched):
                ws.append([label, text])
                review.append(ws.max_row)
    _style_table(ws, 2, ws.max_row, [34, 48])
    for r in review:
        for c in (1, 2):
            ws.cell(r, c).fill = REVIEW_FILL
            ws.cell(r, c).font = Font(color="9A2C1C")
    ws.append([])
    ws.append(["※ 검토 필요 경기: 양 팀 실력 차이가 큰 경기예요. 대진표 시트에 빨간색으로 표시했어요."])

    # 4. 참가자 명단
    ws = wb.create_sheet("참가자 명단")
    ws.append(COLUMNS)
    for p in sched.players:
        ws.append([p.name, p.age, p.gender, p.grade, p.partner])
    _style_table(ws, len(COLUMNS), len(sched.players) + 1)

    wb.save(path)
